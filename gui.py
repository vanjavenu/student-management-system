
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import sqlite3
import os
import shutil
from datetime import date
from openpyxl import Workbook
from PIL import Image, ImageTk, ImageOps, ImageDraw, ImageFont


# =========================================================
# SETTINGS
# =========================================================

DB_NAME = "students.db"
PHOTO_FOLDER = "student_photos"
ID_CARD_FOLDER = "student_id_cards"

os.makedirs(PHOTO_FOLDER, exist_ok=True)
os.makedirs(ID_CARD_FOLDER, exist_ok=True)


# =========================================================
# DATABASE
# =========================================================

def create_tables():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            age TEXT NOT NULL,
            email TEXT NOT NULL,
            phone TEXT NOT NULL,
            course TEXT NOT NULL
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id INTEGER NOT NULL,
            date TEXT NOT NULL,
            status TEXT NOT NULL
        )
    """)

    # Add photo_path if it does not already exist
    cursor.execute("PRAGMA table_info(students)")
    columns = [row[1] for row in cursor.fetchall()]

    if "photo_path" not in columns:
        cursor.execute(
            "ALTER TABLE students ADD COLUMN photo_path TEXT"
        )

    conn.commit()
    conn.close()


create_tables()


# =========================================================
# MAIN WINDOW
# =========================================================

window = tk.Tk()
window.title("Student Management System")
window.geometry("1250x800")
window.minsize(1000, 700)
window.configure(bg="#f3f4f6")


# =========================================================
# GLOBAL VARIABLES
# =========================================================

selected_photo_path = ""

login_window = None


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def get_font(size=12, bold=False):
    if bold:
        return ("Arial", size, "bold")
    return ("Arial", size)


def get_selected_student_id():
    selected = tree.selection()

    if not selected:
        messagebox.showwarning(
            "No Selection",
            "Please select a student first."
        )
        return None

    values = tree.item(selected[0], "values")

    if not values:
        return None

    return int(values[0])


def get_students():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id, name, age, email, phone, course
        FROM students
        ORDER BY id
    """)

    rows = cursor.fetchall()

    conn.close()

    return rows


# =========================================================
# PHOTO FUNCTIONS
# =========================================================

def choose_student_photo():
    global selected_photo_path

    file_path = filedialog.askopenfilename(
        title="Select Student Photo",
        filetypes=[
            ("Image Files", "*.jpg *.jpeg *.png *.webp"),
            ("JPG Files", "*.jpg *.jpeg"),
            ("PNG Files", "*.png"),
            ("All Files", "*.*")
        ]
    )

    if file_path:
        selected_photo_path = file_path
        photo_status_label.config(
            text=os.path.basename(file_path),
            fg="#2563eb"
        )


def clear_photo_selection():
    global selected_photo_path

    selected_photo_path = ""

    photo_status_label.config(
        text="No photo selected",
        fg="#6b7280"
    )


def delete_student_photo_file(photo_path):
    if not photo_path:
        return

    try:
        photo_folder = os.path.abspath(PHOTO_FOLDER)
        photo_file = os.path.abspath(photo_path)

        if (
            photo_file.startswith(photo_folder)
            and os.path.exists(photo_file)
        ):
            os.remove(photo_file)

    except Exception:
        pass


def change_student_photo(student_id, profile_window=None):
    file_path = filedialog.askopenfilename(
        title="Select Student Photo",
        filetypes=[
            ("Image Files", "*.jpg *.jpeg *.png *.webp"),
            ("JPG Files", "*.jpg *.jpeg"),
            ("PNG Files", "*.png"),
            ("All Files", "*.*")
        ]
    )

    if not file_path:
        return

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        "SELECT photo_path FROM students WHERE id = ?",
        (student_id,)
    )

    result = cursor.fetchone()

    if not result:
        conn.close()

        messagebox.showerror(
            "Error",
            "Student not found."
        )

        return

    old_photo_path = result[0]

    try:
        extension = os.path.splitext(file_path)[1].lower()

        new_photo_path = os.path.join(
            PHOTO_FOLDER,
            f"student_{student_id}{extension}"
        )

        # Delete old photo if extension is different
        if old_photo_path:
            old_full = os.path.abspath(old_photo_path)
            new_full = os.path.abspath(new_photo_path)

            if old_full != new_full:
                delete_student_photo_file(old_photo_path)

        shutil.copy2(
            file_path,
            new_photo_path
        )

        cursor.execute(
            """
            UPDATE students
            SET photo_path = ?
            WHERE id = ?
            """,
            (new_photo_path, student_id)
        )

        conn.commit()
        conn.close()

        messagebox.showinfo(
            "Success",
            "Student photo updated successfully!"
        )

        if profile_window is not None:
            profile_window.destroy()

        view_student_profile(student_id)

    except Exception as e:

        conn.close()

        messagebox.showerror(
            "Error",
            f"Could not update photo.\n\n{e}"
        )


# =========================================================
# DASHBOARD
# =========================================================

def get_dashboard_data():

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        "SELECT COUNT(*) FROM students"
    )
    total_students = cursor.fetchone()[0]

    cursor.execute(
        "SELECT COUNT(DISTINCT course) FROM students"
    )
    total_courses = cursor.fetchone()[0]

    today_date = date.today().isoformat()

    cursor.execute(
        """
        SELECT COUNT(*)
        FROM attendance
        WHERE date = ?
        AND status = 'Present'
        """,
        (today_date,)
    )

    present_today = cursor.fetchone()[0]

    cursor.execute(
        """
        SELECT COUNT(*)
        FROM attendance
        WHERE date = ?
        AND status = 'Absent'
        """,
        (today_date,)
    )

    absent_today = cursor.fetchone()[0]

    marked_today = present_today + absent_today

    if marked_today > 0:
        attendance_percentage = (
            present_today / marked_today
        ) * 100
    else:
        attendance_percentage = 0

    conn.close()

    return (
        total_students,
        total_courses,
        present_today,
        absent_today,
        attendance_percentage
    )


def update_dashboard():

    total_students, total_courses, present_today, absent_today, percentage = (
        get_dashboard_data()
    )

    total_students_value.config(
        text=str(total_students)
    )

    total_courses_value.config(
        text=str(total_courses)
    )

    present_today_value.config(
        text=str(present_today)
    )

    absent_today_value.config(
        text=str(absent_today)
    )

    attendance_value.config(
        text=f"{percentage:.1f}%"
    )


# =========================================================
# CLEAR FORM
# =========================================================

def clear_fields():

    global selected_photo_path

    name_entry.delete(0, tk.END)
    age_entry.delete(0, tk.END)
    email_entry.delete(0, tk.END)
    phone_entry.delete(0, tk.END)
    course_entry.delete(0, tk.END)

    selected_photo_path = ""

    photo_status_label.config(
        text="No photo selected",
        fg="#6b7280"
    )

    name_entry.focus()


# =========================================================
# ADD STUDENT
# =========================================================

def add_student():

    global selected_photo_path

    name = name_entry.get().strip()
    age = age_entry.get().strip()
    email = email_entry.get().strip()
    phone = phone_entry.get().strip()
    course = course_entry.get().strip()

    # Validation
    if not name or not age or not email or not phone or not course:

        messagebox.showwarning(
            "Validation",
            "Please fill all fields."
        )

        return

    if not age.isdigit():

        messagebox.showwarning(
            "Validation",
            "Age must contain numbers only."
        )

        return

    if not phone.isdigit():

        messagebox.showwarning(
            "Validation",
            "Phone must contain numbers only."
        )

        return

    if "@" not in email:

        messagebox.showwarning(
            "Validation",
            "Please enter a valid email."
        )

        return

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        INSERT INTO students
        (name, age, email, phone, course, photo_path)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            name,
            age,
            email,
            phone,
            course,
            ""
        )
    )

    student_id = cursor.lastrowid

    # Save photo
    if selected_photo_path:

        try:

            extension = os.path.splitext(
                selected_photo_path
            )[1].lower()

            new_photo_path = os.path.join(
                PHOTO_FOLDER,
                f"student_{student_id}{extension}"
            )

            shutil.copy2(
                selected_photo_path,
                new_photo_path
            )

            cursor.execute(
                """
                UPDATE students
                SET photo_path = ?
                WHERE id = ?
                """,
                (
                    new_photo_path,
                    student_id
                )
            )

        except Exception as e:

            conn.rollback()
            conn.close()

            messagebox.showerror(
                "Photo Error",
                f"Could not save photo.\n\n{e}"
            )

            return

    conn.commit()
    conn.close()

    messagebox.showinfo(
        "Success",
        "Student added successfully!"
    )

    clear_fields()
    load_students()
    update_dashboard()


# =========================================================
# LOAD STUDENTS
# =========================================================

def load_students():

    for item in tree.get_children():
        tree.delete(item)

    rows = get_students()

    for row in rows:

        tree.insert(
            "",
            tk.END,
            values=row
        )


# =========================================================
# SEARCH BY ID
# =========================================================

def search_student_by_id():

    student_id = search_id_entry.get().strip()

    if not student_id.isdigit():

        messagebox.showwarning(
            "Search",
            "Please enter a valid Student ID."
        )

        return

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT id, name, age, email, phone, course
        FROM students
        WHERE id = ?
        """,
        (student_id,)
    )

    student = cursor.fetchone()

    conn.close()

    if not student:

        messagebox.showinfo(
            "Search",
            "Student not found."
        )

        return

    for item in tree.get_children():
        tree.delete(item)

    tree.insert(
        "",
        tk.END,
        values=student
    )


# =========================================================
# SEARCH BY NAME
# =========================================================

def search_student_by_name(*args):

    search_text = name_search_var.get().strip()

    for item in tree.get_children():
        tree.delete(item)

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT id, name, age, email, phone, course
        FROM students
        WHERE name LIKE ?
        ORDER BY id
        """,
        (f"%{search_text}%",)
    )

    rows = cursor.fetchall()

    conn.close()

    for row in rows:

        tree.insert(
            "",
            tk.END,
            values=row
        )


# =========================================================
# COURSE FILTER
# =========================================================

def filter_by_course(*args):

    selected_course = course_filter_var.get()

    for item in tree.get_children():
        tree.delete(item)

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    if selected_course == "All Courses":

        cursor.execute(
            """
            SELECT id, name, age, email, phone, course
            FROM students
            ORDER BY id
            """
        )

    else:

        cursor.execute(
            """
            SELECT id, name, age, email, phone, course
            FROM students
            WHERE course = ?
            ORDER BY id
            """,
            (selected_course,)
        )

    rows = cursor.fetchall()

    conn.close()

    for row in rows:

        tree.insert(
            "",
            tk.END,
            values=row
        )


def load_course_filter():

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT DISTINCT course
        FROM students
        ORDER BY course
        """
    )

    courses = [
        row[0]
        for row in cursor.fetchall()
    ]

    conn.close()

    course_filter["values"] = (
        ["All Courses"] + courses
    )

    if course_filter_var.get() not in course_filter["values"]:

        course_filter_var.set(
            "All Courses"
        )


# =========================================================
# REFRESH
# =========================================================

def refresh_students():

    load_students()
    load_course_filter()
    update_dashboard()


# =========================================================
# EDIT STUDENT
# =========================================================

def edit_selected_student():

    student_id = get_selected_student_id()

    if student_id is None:
        return

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT id, name, age, email, phone, course
        FROM students
        WHERE id = ?
        """,
        (student_id,)
    )

    student = cursor.fetchone()

    conn.close()

    if not student:

        messagebox.showerror(
            "Error",
            "Student not found."
        )

        return

    edit_window = tk.Toplevel(window)
    edit_window.title("Edit Student")
    edit_window.geometry("450x500")
    edit_window.configure(bg="#f3f4f6")
    edit_window.transient(window)
    edit_window.grab_set()

    tk.Label(
        edit_window,
        text="Edit Student",
        font=("Arial", 22, "bold"),
        bg="#f3f4f6",
        fg="#1f2937"
    ).pack(pady=20)

    form = tk.Frame(
        edit_window,
        bg="#f3f4f6"
    )

    form.pack(
        padx=30,
        fill="x"
    )

    tk.Label(
        form,
        text="Name",
        font=get_font(11, True),
        bg="#f3f4f6"
    ).pack(anchor="w")

    edit_name = tk.Entry(
        form,
        font=get_font(12)
    )

    edit_name.pack(
        fill="x",
        pady=(5, 15)
    )

    edit_name.insert(0, student[1])

    tk.Label(
        form,
        text="Age",
        font=get_font(11, True),
        bg="#f3f4f6"
    ).pack(anchor="w")

    edit_age = tk.Entry(
        form,
        font=get_font(12)
    )

    edit_age.pack(
        fill="x",
        pady=(5, 15)
    )

    edit_age.insert(0, student[2])

    tk.Label(
        form,
        text="Email",
        font=get_font(11, True),
        bg="#f3f4f6"
    ).pack(anchor="w")

    edit_email = tk.Entry(
        form,
        font=get_font(12)
    )

    edit_email.pack(
        fill="x",
        pady=(5, 15)
    )

    edit_email.insert(0, student[3])

    tk.Label(
        form,
        text="Phone",
        font=get_font(11, True),
        bg="#f3f4f6"
    ).pack(anchor="w")

    edit_phone = tk.Entry(
        form,
        font=get_font(12)
    )

    edit_phone.pack(
        fill="x",
        pady=(5, 15)
    )

    edit_phone.insert(0, student[4])

    tk.Label(
        form,
        text="Course",
        font=get_font(11, True),
        bg="#f3f4f6"
    ).pack(anchor="w")

    edit_course = tk.Entry(
        form,
        font=get_font(12)
    )

    edit_course.pack(
        fill="x",
        pady=(5, 15)
    )

    edit_course.insert(0, student[5])

    def save_changes():

        name = edit_name.get().strip()
        age = edit_age.get().strip()
        email = edit_email.get().strip()
        phone = edit_phone.get().strip()
        course = edit_course.get().strip()

        if not all(
            [name, age, email, phone, course]
        ):

            messagebox.showwarning(
                "Validation",
                "Please fill all fields."
            )

            return

        if not age.isdigit():

            messagebox.showwarning(
                "Validation",
                "Age must contain numbers only."
            )

            return

        if not phone.isdigit():

            messagebox.showwarning(
                "Validation",
                "Phone must contain numbers only."
            )

            return

        if "@" not in email:

            messagebox.showwarning(
                "Validation",
                "Please enter a valid email."
            )

            return

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        cursor.execute(
            """
            UPDATE students
            SET name = ?,
                age = ?,
                email = ?,
                phone = ?,
                course = ?
            WHERE id = ?
            """,
            (
                name,
                age,
                email,
                phone,
                course,
                student_id
            )
        )

        conn.commit()
        conn.close()

        messagebox.showinfo(
            "Success",
            "Student updated successfully!"
        )

        edit_window.destroy()

        refresh_students()

    tk.Button(
        edit_window,
        text="Save Changes",
        font=get_font(11, True),
        bg="#2563eb",
        fg="white",
        padx=20,
        pady=10,
        command=save_changes
    ).pack(pady=20)


# =========================================================
# DELETE STUDENT
# =========================================================

def delete_selected_student():

    student_id = get_selected_student_id()

    if student_id is None:
        return

    confirm = messagebox.askyesno(
        "Delete Student",
        f"Are you sure you want to delete Student ID {student_id}?"
    )

    if not confirm:
        return

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        "SELECT photo_path FROM students WHERE id = ?",
        (student_id,)
    )

    result = cursor.fetchone()

    photo_path = result[0] if result else ""

    cursor.execute(
        "DELETE FROM attendance WHERE student_id = ?",
        (student_id,)
    )

    cursor.execute(
        "DELETE FROM students WHERE id = ?",
        (student_id,)
    )

    conn.commit()
    conn.close()

    delete_student_photo_file(photo_path)

    messagebox.showinfo(
        "Success",
        "Student deleted successfully!"
    )

    refresh_students()


# =========================================================
# SORT TABLE
# =========================================================

def sort_table(column, reverse=False):

    data = [
        (
            tree.set(child, column),
            child
        )
        for child in tree.get_children("")
    ]

    def sort_key(item):

        value = item[0]

        if column in ["ID", "Age"]:

            try:
                return int(value)
            except:
                return 0

        return value.lower()

    data.sort(
        key=sort_key,
        reverse=reverse
    )

    for index, (_, child) in enumerate(data):

        tree.move(
            child,
            "",
            index
        )

    tree.heading(
        column,
        command=lambda: sort_table(
            column,
            not reverse
        )
    )


# =========================================================
# STUDENT PROFILE
# =========================================================

def view_student_profile(student_id=None):

    if student_id is None:

        student_id = get_selected_student_id()

        if student_id is None:
            return

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT id, name, age, email, phone, course, photo_path
        FROM students
        WHERE id = ?
        """,
        (student_id,)
    )

    student = cursor.fetchone()

    if not student:

        conn.close()

        messagebox.showerror(
            "Error",
            "Student not found."
        )

        return

    cursor.execute(
        """
        SELECT
            COUNT(*),
            SUM(CASE WHEN status = 'Present' THEN 1 ELSE 0 END),
            SUM(CASE WHEN status = 'Absent' THEN 1 ELSE 0 END)
        FROM attendance
        WHERE student_id = ?
        """,
        (student_id,)
    )

    attendance_summary = cursor.fetchone()

    total_days = attendance_summary[0] or 0
    present_days = attendance_summary[1] or 0
    absent_days = attendance_summary[2] or 0

    if total_days > 0:
        attendance_percentage = (
            present_days / total_days
        ) * 100
    else:
        attendance_percentage = 0

    cursor.execute(
        """
        SELECT date, status
        FROM attendance
        WHERE student_id = ?
        ORDER BY date DESC
        """,
        (student_id,)
    )

    attendance_history = cursor.fetchall()

    conn.close()

    # Profile window
    profile_window = tk.Toplevel(window)

    profile_window.title(
        f"Student Profile - {student[1]}"
    )

    profile_window.geometry(
        "800x760"
    )

    profile_window.configure(
        bg="#f3f4f6"
    )

    profile_window.transient(window)

    # Title
    tk.Label(
        profile_window,
        text="Student Profile",
        font=("Arial", 28, "bold"),
        bg="#f3f4f6",
        fg="#1f3b56"
    ).pack(pady=(20, 15))

    # Photo section
    photo_frame = tk.LabelFrame(
        profile_window,
        text="Profile Picture",
        font=("Arial", 12, "bold"),
        bg="#f3f4f6",
        padx=15,
        pady=15
    )

    photo_frame.pack(
        padx=30,
        pady=5
    )

    photo_label = tk.Label(
        photo_frame,
        text="No Photo",
        font=("Arial", 16),
        bg="#eeeeee",
        width=25,
        height=12
    )

    photo_label.pack()

    # Display photo
    if student[6] and os.path.exists(student[6]):

        try:

            image = Image.open(
                student[6]
            ).convert("RGB")

            image = ImageOps.contain(
                image,
                (260, 260),
                method=Image.Resampling.LANCZOS
            )

            photo_image = ImageTk.PhotoImage(
                image
            )

            photo_label.config(
                image=photo_image,
                text=""
            )

            photo_label.image = photo_image

        except Exception:
            pass

    # CHANGE PHOTO BUTTON
    tk.Button(
        photo_frame,
        text="📷 Change / Upload Photo",
        font=("Arial", 11, "bold"),
        bg="#2563eb",
        fg="white",
        padx=15,
        pady=8,
        cursor="hand2",
        command=lambda: change_student_photo(
            student[0],
            profile_window
        )
    ).pack(
        pady=(12, 0)
    )

    # Information frame
    info_frame = tk.LabelFrame(
        profile_window,
        text="Student Information",
        font=("Arial", 12, "bold"),
        bg="white",
        padx=15,
        pady=15
    )

    info_frame.pack(
        fill="x",
        padx=30,
        pady=15
    )

    info_data = [
        ("Student ID:", student[0]),
        ("Name:", student[1]),
        ("Age:", student[2]),
        ("Email:", student[3]),
        ("Phone:", student[4]),
        ("Course:", student[5]),
        ("Total Attendance Days:", total_days),
        ("Present Days:", present_days),
        ("Absent Days:", absent_days),
        (
            "Attendance:",
            f"{attendance_percentage:.1f}%"
        )
    ]

    for row_index, (label_text, value) in enumerate(info_data):

        column = 0 if row_index < 5 else 2
        row = row_index if row_index < 5 else row_index - 5

        tk.Label(
            info_frame,
            text=label_text,
            font=("Arial", 11, "bold"),
            bg="white",
            anchor="w"
        ).grid(
            row=row,
            column=column,
            sticky="w",
            padx=5,
            pady=8
        )

        tk.Label(
            info_frame,
            text=value,
            font=("Arial", 11),
            bg="white",
            anchor="w"
        ).grid(
            row=row,
            column=column + 1,
            sticky="w",
            padx=25,
            pady=8
        )

    info_frame.columnconfigure(
        1,
        weight=1
    )

    info_frame.columnconfigure(
        3,
        weight=1
    )

    # Attendance History
    history_frame = tk.LabelFrame(
        profile_window,
        text="Attendance History",
        font=("Arial", 12, "bold"),
        bg="#f3f4f6",
        padx=8,
        pady=8
    )

    history_frame.pack(
        fill="both",
        expand=True,
        padx=30,
        pady=(0, 20)
    )

    history_tree = ttk.Treeview(
        history_frame,
        columns=("Date", "Status"),
        show="headings",
        height=7
    )

    history_tree.heading(
        "Date",
        text="Date"
    )

    history_tree.heading(
        "Status",
        text="Status"
    )

    history_tree.column(
        "Date",
        width=250,
        anchor="center"
    )

    history_tree.column(
        "Status",
        width=250,
        anchor="center"
    )

    history_tree.pack(
        fill="both",
        expand=True
    )

    for record in attendance_history:

        history_tree.insert(
            "",
            tk.END,
            values=record
        )

    # Generate ID Card button
    tk.Button(
        profile_window,
        text="🎫 Generate ID Card",
        font=("Arial", 11, "bold"),
        bg="#059669",
        fg="white",
        padx=20,
        pady=8,
        command=lambda: generate_id_card(student[0])
    ).pack(
        pady=(0, 15)
    )


# =========================================================
# ATTENDANCE
# =========================================================

def mark_attendance(
    student_id,
    attendance_date,
    status
):

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT id
        FROM attendance
        WHERE student_id = ?
        AND date = ?
        """,
        (
            student_id,
            attendance_date
        )
    )

    existing = cursor.fetchone()

    if existing:

        cursor.execute(
            """
            UPDATE attendance
            SET status = ?
            WHERE student_id = ?
            AND date = ?
            """,
            (
                status,
                student_id,
                attendance_date
            )
        )

    else:

        cursor.execute(
            """
            INSERT INTO attendance
            (student_id, date, status)
            VALUES (?, ?, ?)
            """,
            (
                student_id,
                attendance_date,
                status
            )
        )

    conn.commit()
    conn.close()


def open_attendance_window():

    attendance_window = tk.Toplevel(window)

    attendance_window.title(
        "Attendance Management"
    )

    attendance_window.geometry(
        "850x650"
    )

    attendance_window.configure(
        bg="#f3f4f6"
    )

    tk.Label(
        attendance_window,
        text="Attendance Management",
        font=("Arial", 24, "bold"),
        bg="#f3f4f6",
        fg="#1f3b56"
    ).pack(pady=20)

    control_frame = tk.Frame(
        attendance_window,
        bg="#f3f4f6"
    )

    control_frame.pack(
        fill="x",
        padx=30
    )

    # Student
    tk.Label(
        control_frame,
        text="Student:",
        font=get_font(11, True),
        bg="#f3f4f6"
    ).grid(
        row=0,
        column=0,
        padx=5,
        pady=10
    )

    student_var = tk.StringVar()

    student_combo = ttk.Combobox(
        control_frame,
        textvariable=student_var,
        state="readonly",
        width=35
    )

    student_combo.grid(
        row=0,
        column=1,
        padx=5
    )

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT id, name
        FROM students
        ORDER BY id
        """
    )

    student_list = cursor.fetchall()

    conn.close()

    student_combo["values"] = [
        f"{row[0]} - {row[1]}"
        for row in student_list
    ]

    # Date
    tk.Label(
        control_frame,
        text="Date:",
        font=get_font(11, True),
        bg="#f3f4f6"
    ).grid(
        row=1,
        column=0,
        padx=5,
        pady=10
    )

    date_entry = tk.Entry(
        control_frame,
        font=get_font(11),
        width=25
    )

    date_entry.grid(
        row=1,
        column=1,
        padx=5,
        sticky="w"
    )

    date_entry.insert(
        0,
        date.today().isoformat()
    )

    # Attendance buttons
    def mark_present():

        selected = student_var.get()

        if not selected:

            messagebox.showwarning(
                "Attendance",
                "Please select a student."
            )

            return

        student_id = int(
            selected.split(" - ")[0]
        )

        attendance_date = date_entry.get().strip()

        if not attendance_date:

            messagebox.showwarning(
                "Attendance",
                "Please enter a date."
            )

            return

        mark_attendance(
            student_id,
            attendance_date,
            "Present"
        )

        load_attendance()

        update_dashboard()

        messagebox.showinfo(
            "Success",
            "Present marked successfully!"
        )

    def mark_absent():

        selected = student_var.get()

        if not selected:

            messagebox.showwarning(
                "Attendance",
                "Please select a student."
            )

            return

        student_id = int(
            selected.split(" - ")[0]
        )

        attendance_date = date_entry.get().strip()

        if not attendance_date:

            messagebox.showwarning(
                "Attendance",
                "Please enter a date."
            )

            return

        mark_attendance(
            student_id,
            attendance_date,
            "Absent"
        )

        load_attendance()

        update_dashboard()

        messagebox.showinfo(
            "Success",
            "Absent marked successfully!"
        )

    button_frame = tk.Frame(
        attendance_window,
        bg="#f3f4f6"
    )

    button_frame.pack(
        pady=15
    )

    tk.Button(
        button_frame,
        text="✓ Present",
        font=get_font(11, True),
        bg="#16a34a",
        fg="white",
        padx=20,
        pady=8,
        command=mark_present
    ).pack(
        side="left",
        padx=8
    )

    tk.Button(
        button_frame,
        text="✗ Absent",
        font=get_font(11, True),
        bg="#dc2626",
        fg="white",
        padx=20,
        pady=8,
        command=mark_absent
    ).pack(
        side="left",
        padx=8
    )

    # Attendance table
    table_frame = tk.Frame(
        attendance_window,
        bg="#f3f4f6"
    )

    table_frame.pack(
        fill="both",
        expand=True,
        padx=30,
        pady=10
    )

    attendance_tree = ttk.Treeview(
        table_frame,
        columns=(
            "ID",
            "Name",
            "Date",
            "Status"
        ),
        show="headings"
    )

    for column in (
        "ID",
        "Name",
        "Date",
        "Status"
    ):

        attendance_tree.heading(
            column,
            text=column
        )

    attendance_tree.column(
        "ID",
        width=70,
        anchor="center"
    )

    attendance_tree.column(
        "Name",
        width=200
    )

    attendance_tree.column(
        "Date",
        width=180,
        anchor="center"
    )

    attendance_tree.column(
        "Status",
        width=150,
        anchor="center"
    )

    attendance_tree.pack(
        fill="both",
        expand=True
    )

    def load_attendance():

        for item in attendance_tree.get_children():
            attendance_tree.delete(item)

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT
                students.id,
                students.name,
                attendance.date,
                attendance.status
            FROM attendance
            JOIN students
            ON students.id = attendance.student_id
            ORDER BY attendance.date DESC
            """
        )

        rows = cursor.fetchall()

        conn.close()

        for row in rows:

            attendance_tree.insert(
                "",
                tk.END,
                values=row
            )

    load_attendance()

    bottom_frame = tk.Frame(
        attendance_window,
        bg="#f3f4f6"
    )

    bottom_frame.pack(
        pady=15
    )

    tk.Button(
        bottom_frame,
        text="Attendance Summary",
        font=get_font(10, True),
        bg="#7c3aed",
        fg="white",
        padx=15,
        pady=7,
        command=open_attendance_summary
    ).pack(
        side="left",
        padx=5
    )

    tk.Button(
        bottom_frame,
        text="Export Attendance",
        font=get_font(10, True),
        bg="#0284c7",
        fg="white",
        padx=15,
        pady=7,
        command=export_attendance
    ).pack(
        side="left",
        padx=5
    )


# =========================================================
# ATTENDANCE SUMMARY
# =========================================================

def open_attendance_summary():

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            students.id,
            students.name,
            students.course,
            COUNT(attendance.id),
            SUM(
                CASE
                    WHEN attendance.status = 'Present'
                    THEN 1 ELSE 0
                END
            ),
            SUM(
                CASE
                    WHEN attendance.status = 'Absent'
                    THEN 1 ELSE 0
                END
            )
        FROM students
        LEFT JOIN attendance
        ON students.id = attendance.student_id
        GROUP BY students.id
        ORDER BY students.id
        """
    )

    rows = cursor.fetchall()

    conn.close()

    summary_window = tk.Toplevel(window)

    summary_window.title(
        "Attendance Summary"
    )

    summary_window.geometry(
        "850x550"
    )

    summary_window.configure(
        bg="#f3f4f6"
    )

    tk.Label(
        summary_window,
        text="Attendance Summary",
        font=("Arial", 24, "bold"),
        bg="#f3f4f6",
        fg="#1f3b56"
    ).pack(pady=20)

    summary_tree = ttk.Treeview(
        summary_window,
        columns=(
            "ID",
            "Name",
            "Course",
            "Total",
            "Present",
            "Absent",
            "Percentage"
        ),
        show="headings"
    )

    for column in (
        "ID",
        "Name",
        "Course",
        "Total",
        "Present",
        "Absent",
        "Percentage"
    ):

        summary_tree.heading(
            column,
            text=column
        )

    for column in (
        "ID",
        "Total",
        "Present",
        "Absent"
    ):

        summary_tree.column(
            column,
            width=80,
            anchor="center"
        )

    summary_tree.column(
        "Name",
        width=180
    )

    summary_tree.column(
        "Course",
        width=130
    )

    summary_tree.column(
        "Percentage",
        width=110,
        anchor="center"
    )

    summary_tree.pack(
        fill="both",
        expand=True,
        padx=25,
        pady=15
    )

    for row in rows:

        total = row[3] or 0
        present = row[4] or 0
        absent = row[5] or 0

        percentage = (
            (present / total) * 100
            if total > 0
            else 0
        )

        summary_tree.insert(
            "",
            tk.END,
            values=(
                row[0],
                row[1],
                row[2],
                total,
                present,
                absent,
                f"{percentage:.1f}%"
            )
        )


# =========================================================
# EXPORT ATTENDANCE
# =========================================================

def export_attendance():

    file_path = filedialog.asksaveasfilename(
        title="Save Attendance",
        defaultextension=".xlsx",
        filetypes=[
            ("Excel Files", "*.xlsx")
        ],
        initialfile="attendance.xlsx"
    )

    if not file_path:
        return

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            students.id,
            students.name,
            students.course,
            attendance.date,
            attendance.status
        FROM attendance
        JOIN students
        ON students.id = attendance.student_id
        ORDER BY attendance.date DESC
        """
    )

    rows = cursor.fetchall()

    conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"

    headers = [
        "Student ID",
        "Name",
        "Course",
        "Date",
        "Status"
    ]

    sheet.append(headers)

    for row in rows:
        sheet.append(row)

    workbook.save(file_path)

    messagebox.showinfo(
        "Export Complete",
        f"Attendance exported successfully!\n\n{file_path}"
    )


# =========================================================
# MONTHLY REPORT
# =========================================================

def get_monthly_report(month, year):

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    month_text = f"{year:04d}-{month:02d}"

    cursor.execute(
        """
        SELECT
            students.id,
            students.name,
            students.course,
            COUNT(attendance.id),
            SUM(
                CASE
                    WHEN attendance.status = 'Present'
                    THEN 1 ELSE 0
                END
            ),
            SUM(
                CASE
                    WHEN attendance.status = 'Absent'
                    THEN 1 ELSE 0
                END
            )
        FROM students
        LEFT JOIN attendance
        ON students.id = attendance.student_id
        AND attendance.date LIKE ?
        GROUP BY students.id
        ORDER BY students.id
        """,
        (f"{month_text}%",)
    )

    rows = cursor.fetchall()

    conn.close()

    report = []

    for row in rows:

        total = row[3] or 0
        present = row[4] or 0
        absent = row[5] or 0

        percentage = (
            (present / total) * 100
            if total > 0
            else 0
        )

        report.append(
            (
                row[0],
                row[1],
                row[2],
                total,
                present,
                absent,
                f"{percentage:.1f}%"
            )
        )

    return report


def open_monthly_report():

    report_window = tk.Toplevel(window)

    report_window.title(
        "Monthly Attendance Report"
    )

    report_window.geometry(
        "900x600"
    )

    report_window.configure(
        bg="#f3f4f6"
    )

    tk.Label(
        report_window,
        text="Monthly Attendance Report",
        font=("Arial", 24, "bold"),
        bg="#f3f4f6",
        fg="#1f3b56"
    ).pack(pady=20)

    controls = tk.Frame(
        report_window,
        bg="#f3f4f6"
    )

    controls.pack()

    tk.Label(
        controls,
        text="Month:",
        font=get_font(11, True),
        bg="#f3f4f6"
    ).grid(
        row=0,
        column=0,
        padx=5
    )

    month_var = tk.StringVar(
        value=str(date.today().month)
    )

    month_combo = ttk.Combobox(
        controls,
        textvariable=month_var,
        values=[
            str(i)
            for i in range(1, 13)
        ],
        width=8,
        state="readonly"
    )

    month_combo.grid(
        row=0,
        column=1,
        padx=5
    )

    tk.Label(
        controls,
        text="Year:",
        font=get_font(11, True),
        bg="#f3f4f6"
    ).grid(
        row=0,
        column=2,
        padx=5
    )

    year_var = tk.StringVar(
        value=str(date.today().year)
    )

    year_entry = tk.Entry(
        controls,
        textvariable=year_var,
        width=10,
        font=get_font(11)
    )

    year_entry.grid(
        row=0,
        column=3,
        padx=5
    )

    table = ttk.Treeview(
        report_window,
        columns=(
            "ID",
            "Name",
            "Course",
            "Total",
            "Present",
            "Absent",
            "Percentage"
        ),
        show="headings"
    )

    for column in (
        "ID",
        "Name",
        "Course",
        "Total",
        "Present",
        "Absent",
        "Percentage"
    ):

        table.heading(
            column,
            text=column
        )

    table.pack(
        fill="both",
        expand=True,
        padx=25,
        pady=20
    )

    current_report = []

    def generate_report():

        try:
            month = int(
                month_var.get()
            )

            year = int(
                year_var.get()
            )

        except ValueError:

            messagebox.showwarning(
                "Report",
                "Please enter a valid month and year."
            )

            return

        current_report.clear()

        rows = get_monthly_report(
            month,
            year
        )

        current_report.extend(rows)

        for item in table.get_children():
            table.delete(item)

        for row in rows:

            table.insert(
                "",
                tk.END,
                values=row
            )

    def export_report():

        if not current_report:

            messagebox.showwarning(
                "Export",
                "Please generate the report first."
            )

            return

        month = int(
            month_var.get()
        )

        year = int(
            year_var.get()
        )

        file_path = filedialog.asksaveasfilename(
            title="Save Monthly Report",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel Files", "*.xlsx")
            ],
            initialfile=f"attendance_{year}_{month:02d}.xlsx"
        )

        if not file_path:
            return

        workbook = Workbook()
        sheet = workbook.active

        sheet.title = "Monthly Attendance"

        headers = [
            "Student ID",
            "Name",
            "Course",
            "Total Days",
            "Present",
            "Absent",
            "Attendance %"
        ]

        sheet.append(headers)

        for row in current_report:
            sheet.append(row)

        workbook.save(file_path)

        messagebox.showinfo(
            "Export Complete",
            "Monthly report exported successfully!"
        )

    button_frame = tk.Frame(
        report_window,
        bg="#f3f4f6"
    )

    button_frame.pack(
        pady=10
    )

    tk.Button(
        button_frame,
        text="Generate Report",
        font=get_font(10, True),
        bg="#2563eb",
        fg="white",
        padx=15,
        pady=7,
        command=generate_report
    ).pack(
        side="left",
        padx=5
    )

    tk.Button(
        button_frame,
        text="Export Report",
        font=get_font(10, True),
        bg="#059669",
        fg="white",
        padx=15,
        pady=7,
        command=export_report
    ).pack(
        side="left",
        padx=5
    )


# =========================================================
# EXPORT STUDENTS
# =========================================================

def export_to_excel():

    file_path = filedialog.asksaveasfilename(
        title="Save Students",
        defaultextension=".xlsx",
        filetypes=[
            ("Excel Files", "*.xlsx")
        ],
        initialfile="students.xlsx"
    )

    if not file_path:
        return

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            id,
            name,
            age,
            email,
            phone,
            course
        FROM students
        ORDER BY id
        """
    )

    rows = cursor.fetchall()

    conn.close()

    workbook = Workbook()
    sheet = workbook.active

    sheet.title = "Students"

    headers = [
        "Student ID",
        "Name",
        "Age",
        "Email",
        "Phone",
        "Course"
    ]

    sheet.append(headers)

    for row in rows:
        sheet.append(row)

    workbook.save(file_path)

    messagebox.showinfo(
        "Export Complete",
        "Students exported successfully!"
    )


# =========================================================
# STUDENT ID CARD
# =========================================================

def create_id_card_image(student):

    width = 900
    height = 560

    image = Image.new(
        "RGB",
        (width, height),
        "white"
    )

    draw = ImageDraw.Draw(image)

    # Fonts
    title_font = get_pil_font(
        38,
        bold=True
    )

    subtitle_font = get_pil_font(
        22,
        bold=False
    )

    label_font = get_pil_font(
        24,
        bold=True
    )

    value_font = get_pil_font(
        24,
        bold=False
    )

    footer_font = get_pil_font(
        18,
        bold=False
    )

    # Header
    draw.rectangle(
        (0, 0, width, 140),
        fill="#2f4356"
    )

    draw.text(
        (width // 2, 50),
        "STUDENT MANAGEMENT SYSTEM",
        fill="white",
        font=title_font,
        anchor="mm"
    )

    draw.text(
        (width // 2, 110),
        "STUDENT ID CARD",
        fill="white",
        font=subtitle_font,
        anchor="mm"
    )

    # Main border
    draw.rectangle(
        (10, 10, width - 10, height - 10),
        outline="#2f4356",
        width=5
    )

    # Photo box
    photo_x1 = 60
    photo_y1 = 190
    photo_x2 = 310
    photo_y2 = 450

    draw.rounded_rectangle(
        (
            photo_x1,
            photo_y1,
            photo_x2,
            photo_y2
        ),
        radius=20,
        outline="#2f4356",
        width=5,
        fill="#eeeeee"
    )

    photo_path = student[6]

    if photo_path and os.path.exists(photo_path):

        try:

            student_image = Image.open(
                photo_path
            ).convert("RGB")

            student_image = ImageOps.contain(
                student_image,
                (
                    photo_x2 - photo_x1 - 20,
                    photo_y2 - photo_y1 - 20
                ),
                method=Image.Resampling.LANCZOS
            )

            px = photo_x1 + (
                photo_x2 - photo_x1 -
                student_image.width
            ) // 2

            py = photo_y1 + (
                photo_y2 - photo_y1 -
                student_image.height
            ) // 2

            image.paste(
                student_image,
                (px, py)
            )

        except Exception:

            draw.text(
                (
                    (photo_x1 + photo_x2) // 2,
                    (photo_y1 + photo_y2) // 2
                ),
                "NO PHOTO",
                fill="#777777",
                font=label_font,
                anchor="mm"
            )

    else:

        draw.text(
            (
                (photo_x1 + photo_x2) // 2,
                (photo_y1 + photo_y2) // 2
            ),
            "NO PHOTO",
            fill="#777777",
            font=label_font,
            anchor="mm"
        )

    # Student information
    x_label = 360
    x_value = 560

    info = [
        ("Student ID:", student[0]),
        ("Name:", student[1]),
        ("Age:", student[2]),
        ("Course:", student[5]),
        ("Phone:", student[4]),
        ("Email:", student[3])
    ]

    y = 205

    for label, value in info:

        draw.text(
            (x_label, y),
            label,
            fill="#1f2937",
            font=label_font
        )

        draw.text(
            (x_value, y),
            str(value),
            fill="#374151",
            font=value_font
        )

        y += 50

    # Footer line
    draw.line(
        (60, 485, 840, 485),
        fill="#cccccc",
        width=2
    )

    draw.text(
        (width // 2, 520),
        "This card is issued by Student Management System",
        fill="#999999",
        font=footer_font,
        anchor="mm"
    )

    return image


def get_pil_font(size, bold=False):

    font_paths = []

    if bold:

        font_paths = [
            "C:/Windows/Fonts/arialbd.ttf",
            "C:/Windows/Fonts/calibrib.ttf",
            "C:/Windows/Fonts/segoeuib.ttf"
        ]

    else:

        font_paths = [
            "C:/Windows/Fonts/arial.ttf",
            "C:/Windows/Fonts/calibri.ttf",
            "C:/Windows/Fonts/segoeui.ttf"
        ]

    for path in font_paths:

        if os.path.exists(path):

            try:
                return ImageFont.truetype(
                    path,
                    size
                )
            except:
                pass

    return ImageFont.load_default()


def generate_id_card(student_id=None):

    if student_id is None:

        student_id = get_selected_student_id()

        if student_id is None:
            return

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            id,
            name,
            age,
            email,
            phone,
            course,
            photo_path
        FROM students
        WHERE id = ?
        """,
        (student_id,)
    )

    student = cursor.fetchone()

    conn.close()

    if not student:

        messagebox.showerror(
            "Error",
            "Student not found."
        )

        return

    try:

        card_image = create_id_card_image(
            student
        )

    except Exception as e:

        messagebox.showerror(
            "Error",
            f"Could not create ID card.\n\n{e}"
        )

        return

    preview_window = tk.Toplevel(window)

    preview_window.title(
        f"ID Card - {student[1]}"
    )

    preview_window.geometry(
        "1000x760"
    )

    preview_window.configure(
        bg="#f3f4f6"
    )

    tk.Label(
        preview_window,
        text="Student ID Card",
        font=("Arial", 28, "bold"),
        bg="#f3f4f6",
        fg="#1f3b56"
    ).pack(pady=15)

    preview_image = ImageTk.PhotoImage(
        card_image.resize(
            (900, 560),
            Image.Resampling.LANCZOS
        )
    )

    image_label = tk.Label(
        preview_window,
        image=preview_image,
        bg="#f3f4f6"
    )

    image_label.image = preview_image

    image_label.pack(
        pady=10
    )

    def save_id_card():

        file_path = filedialog.asksaveasfilename(
            title="Save Student ID Card",
            defaultextension=".png",
            filetypes=[
                ("PNG Image", "*.png"),
                ("JPEG Image", "*.jpg")
            ],
            initialfile=f"student_id_{student[0]}.png",
            initialdir=ID_CARD_FOLDER
        )

        if not file_path:
            return

        try:

            card_image.save(
                file_path
            )

            messagebox.showinfo(
                "Success",
                f"ID card saved successfully!\n\n{file_path}"
            )

        except Exception as e:

            messagebox.showerror(
                "Error",
                f"Could not save ID card.\n\n{e}"
            )

    tk.Button(
        preview_window,
        text="💾 Save ID Card",
        font=("Arial", 12, "bold"),
        bg="#2563eb",
        fg="white",
        padx=25,
        pady=10,
        command=save_id_card
    ).pack(
        pady=15
    )


# =========================================================
# LOGOUT
# =========================================================

def logout():

    confirm = messagebox.askyesno(
        "Logout",
        "Are you sure you want to logout?"
    )

    if not confirm:
        return

    window.withdraw()

    if login_window is not None:

        username_login_entry.delete(
            0,
            tk.END
        )

        password_login_entry.delete(
            0,
            tk.END
        )

        password_login_entry.config(
            show="*"
        )

        show_password_var.set(False)

        login_window.deiconify()

        username_login_entry.focus()


# =========================================================
# LOGIN
# =========================================================

def login():

    username = username_login_entry.get().strip()
    password = password_login_entry.get()

    if username == "admin" and password == "1234":

        login_window.withdraw()

        window.deiconify()

        username_login_entry.delete(
            0,
            tk.END
        )

        password_login_entry.delete(
            0,
            tk.END
        )

        name_entry.focus()

    else:

        messagebox.showerror(
            "Login Failed",
            "Invalid username or password."
        )


def toggle_password():

    if show_password_var.get():

        password_login_entry.config(
            show=""
        )

    else:

        password_login_entry.config(
            show="*"
        )


# =========================================================
# SCROLLABLE MAIN AREA
# =========================================================

main_canvas = tk.Canvas(
    window,
    bg="#f3f4f6",
    highlightthickness=0
)

main_scrollbar = ttk.Scrollbar(
    window,
    orient="vertical",
    command=main_canvas.yview
)

main_frame = tk.Frame(
    main_canvas,
    bg="#f3f4f6"
)

main_frame.bind(
    "<Configure>",
    lambda e: main_canvas.configure(
        scrollregion=main_canvas.bbox("all")
    )
)

main_canvas.create_window(
    (0, 0),
    window=main_frame,
    anchor="nw"
)

main_canvas.configure(
    yscrollcommand=main_scrollbar.set
)

main_canvas.pack(
    side="left",
    fill="both",
    expand=True
)

main_scrollbar.pack(
    side="right",
    fill="y"
)


# =========================================================
# HEADER
# =========================================================

header = tk.Frame(
    main_frame,
    bg="#1f3b56",
    height=100
)

header.pack(
    fill="x"
)

tk.Label(
    header,
    text="Student Management System",
    font=("Arial", 28, "bold"),
    bg="#1f3b56",
    fg="white"
).pack(
    pady=(20, 0)
)

tk.Label(
    header,
    text="Manage students, attendance and ID cards",
    font=("Arial", 12),
    bg="#1f3b56",
    fg="#dbeafe"
).pack(
    pady=(5, 15)
)


# =========================================================
# DASHBOARD
# =========================================================

dashboard_frame = tk.Frame(
    main_frame,
    bg="#f3f4f6"
)

dashboard_frame.pack(
    fill="x",
    padx=25,
    pady=20
)


def create_dashboard_card(
    parent,
    title,
    column
):

    card = tk.Frame(
        parent,
        bg="white",
        bd=1,
        relief="solid",
        width=200,
        height=120
    )

    card.grid(
        row=0,
        column=column,
        padx=7,
        sticky="nsew"
    )

    card.grid_propagate(False)

    tk.Label(
        card,
        text=title,
        font=("Arial", 11, "bold"),
        bg="white",
        fg="#6b7280"
    ).pack(
        pady=(18, 5)
    )

    value_label = tk.Label(
        card,
        text="0",
        font=("Arial", 25, "bold"),
        bg="white",
        fg="#1f3b56"
    )

    value_label.pack()

    return value_label


for i in range(5):

    dashboard_frame.columnconfigure(
        i,
        weight=1
    )


total_students_value = create_dashboard_card(
    dashboard_frame,
    "TOTAL STUDENTS",
    0
)

total_courses_value = create_dashboard_card(
    dashboard_frame,
    "TOTAL COURSES",
    1
)

present_today_value = create_dashboard_card(
    dashboard_frame,
    "PRESENT TODAY",
    2
)

absent_today_value = create_dashboard_card(
    dashboard_frame,
    "ABSENT TODAY",
    3
)

attendance_value = create_dashboard_card(
    dashboard_frame,
    "TODAY'S ATTENDANCE",
    4
)


# =========================================================
# ADD STUDENT
# =========================================================

add_frame = tk.LabelFrame(
    main_frame,
    text="Add Student",
    font=("Arial", 13, "bold"),
    bg="#f3f4f6",
    padx=20,
    pady=20
)

add_frame.pack(
    fill="x",
    padx=25,
    pady=10
)


# Row 1
tk.Label(
    add_frame,
    text="Name:",
    font=get_font(11, True),
    bg="#f3f4f6"
).grid(
    row=0,
    column=0,
    sticky="w",
    padx=5,
    pady=7
)

name_entry = tk.Entry(
    add_frame,
    font=get_font(11),
    width=25
)

name_entry.grid(
    row=0,
    column=1,
    padx=5,
    pady=7
)


tk.Label(
    add_frame,
    text="Age:",
    font=get_font(11, True),
    bg="#f3f4f6"
).grid(
    row=0,
    column=2,
    sticky="w",
    padx=5,
    pady=7
)

age_entry = tk.Entry(
    add_frame,
    font=get_font(11),
    width=20
)

age_entry.grid(
    row=0,
    column=3,
    padx=5,
    pady=7
)


# Row 2
tk.Label(
    add_frame,
    text="Email:",
    font=get_font(11, True),
    bg="#f3f4f6"
).grid(
    row=1,
    column=0,
    sticky="w",
    padx=5,
    pady=7
)

email_entry = tk.Entry(
    add_frame,
    font=get_font(11),
    width=25
)

email_entry.grid(
    row=1,
    column=1,
    padx=5,
    pady=7
)


tk.Label(
    add_frame,
    text="Phone:",
    font=get_font(11, True),
    bg="#f3f4f6"
).grid(
    row=1,
    column=2,
    sticky="w",
    padx=5,
    pady=7
)

phone_entry = tk.Entry(
    add_frame,
    font=get_font(11),
    width=20
)

phone_entry.grid(
    row=1,
    column=3,
    padx=5,
    pady=7
)


# Row 3
tk.Label(
    add_frame,
    text="Course:",
    font=get_font(11, True),
    bg="#f3f4f6"
).grid(
    row=2,
    column=0,
    sticky="w",
    padx=5,
    pady=7
)

course_entry = tk.Entry(
    add_frame,
    font=get_font(11),
    width=25
)

course_entry.grid(
    row=2,
    column=1,
    padx=5,
    pady=7
)


# Photo
tk.Label(
    add_frame,
    text="Photo:",
    font=get_font(11, True),
    bg="#f3f4f6"
).grid(
    row=2,
    column=2,
    sticky="w",
    padx=5,
    pady=7
)

photo_button_frame = tk.Frame(
    add_frame,
    bg="#f3f4f6"
)

photo_button_frame.grid(
    row=2,
    column=3,
    sticky="w",
    padx=5,
    pady=7
)

tk.Button(
    photo_button_frame,
    text="Choose Photo",
    font=get_font(10, True),
    bg="#7c3aed",
    fg="white",
    command=choose_student_photo
).pack(
    side="left"
)

tk.Button(
    photo_button_frame,
    text="Clear",
    font=get_font(10),
    command=clear_photo_selection
).pack(
    side="left",
    padx=5
)

photo_status_label = tk.Label(
    add_frame,
    text="No photo selected",
    font=("Arial", 9),
    bg="#f3f4f6",
    fg="#6b7280"
)

photo_status_label.grid(
    row=3,
    column=3,
    sticky="w",
    padx=5
)


# Add/Clear buttons
button_frame = tk.Frame(
    add_frame,
    bg="#f3f4f6"
)

button_frame.grid(
    row=4,
    column=0,
    columnspan=4,
    pady=15
)

tk.Button(
    button_frame,
    text="➕ Add Student",
    font=("Arial", 11, "bold"),
    bg="#16a34a",
    fg="white",
    padx=20,
    pady=8,
    command=add_student
).pack(
    side="left",
    padx=8
)

tk.Button(
    button_frame,
    text="Clear",
    font=("Arial", 11, "bold"),
    bg="#6b7280",
    fg="white",
    padx=20,
    pady=8,
    command=clear_fields
).pack(
    side="left",
    padx=8
)


# =========================================================
# SEARCH / FILTER
# =========================================================

search_frame = tk.LabelFrame(
    main_frame,
    text="Search & Filter",
    font=("Arial", 13, "bold"),
    bg="#f3f4f6",
    padx=15,
    pady=15
)

search_frame.pack(
    fill="x",
    padx=25,
    pady=10
)


# Search ID
tk.Label(
    search_frame,
    text="Search by ID:",
    font=get_font(10, True),
    bg="#f3f4f6"
).grid(
    row=0,
    column=0,
    padx=5
)

search_id_entry = tk.Entry(
    search_frame,
    font=get_font(10),
    width=12
)

search_id_entry.grid(
    row=0,
    column=1,
    padx=5
)

tk.Button(
    search_frame,
    text="Search",
    font=get_font(10, True),
    bg="#2563eb",
    fg="white",
    command=search_student_by_id
).grid(
    row=0,
    column=2,
    padx=5
)


# Search name
tk.Label(
    search_frame,
    text="Search by Name:",
    font=get_font(10, True),
    bg="#f3f4f6"
).grid(
    row=0,
    column=3,
    padx=5
)

name_search_var = tk.StringVar()

name_search_entry = tk.Entry(
    search_frame,
    textvariable=name_search_var,
    font=get_font(10),
    width=18
)

name_search_entry.grid(
    row=0,
    column=4,
    padx=5
)

name_search_var.trace_add(
    "write",
    search_student_by_name
)


# Course filter
tk.Label(
    search_frame,
    text="Course:",
    font=get_font(10, True),
    bg="#f3f4f6"
).grid(
    row=0,
    column=5,
    padx=5
)

course_filter_var = tk.StringVar(
    value="All Courses"
)

course_filter = ttk.Combobox(
    search_frame,
    textvariable=course_filter_var,
    state="readonly",
    width=18
)

course_filter.grid(
    row=0,
    column=6,
    padx=5
)

course_filter.bind(
    "<<ComboboxSelected>>",
    filter_by_course
)


tk.Button(
    search_frame,
    text="Refresh",
    font=get_font(10, True),
    bg="#059669",
    fg="white",
    command=refresh_students
).grid(
    row=0,
    column=7,
    padx=5
)


# =========================================================
# STUDENT TABLE
# =========================================================

table_frame = tk.LabelFrame(
    main_frame,
    text="Students",
    font=("Arial", 13, "bold"),
    bg="#f3f4f6",
    padx=10,
    pady=10
)

table_frame.pack(
    fill="both",
    expand=True,
    padx=25,
    pady=10
)

table_columns = (
    "ID",
    "Name",
    "Age",
    "Email",
    "Phone",
    "Course"
)

tree = ttk.Treeview(
    table_frame,
    columns=table_columns,
    show="headings",
    height=10
)

for column in table_columns:

    tree.heading(
        column,
        text=column,
        command=lambda c=column: sort_table(c)
    )


tree.column(
    "ID",
    width=70,
    anchor="center"
)

tree.column(
    "Name",
    width=150
)

tree.column(
    "Age",
    width=70,
    anchor="center"
)

tree.column(
    "Email",
    width=220
)

tree.column(
    "Phone",
    width=150
)

tree.column(
    "Course",
    width=130
)

tree.pack(
    side="left",
    fill="both",
    expand=True
)


table_scrollbar = ttk.Scrollbar(
    table_frame,
    orient="vertical",
    command=tree.yview
)

table_scrollbar.pack(
    side="right",
    fill="y"
)

tree.configure(
    yscrollcommand=table_scrollbar.set
)


# Double click profile
tree.bind(
    "<Double-1>",
    lambda event: view_student_profile()
)


# =========================================================
# TABLE BUTTONS
# =========================================================

table_button_frame = tk.Frame(
    main_frame,
    bg="#f3f4f6"
)

table_button_frame.pack(
    pady=10
)

tk.Button(
    table_button_frame,
    text="👤 View Profile",
    font=get_font(10, True),
    bg="#7c3aed",
    fg="white",
    padx=15,
    pady=7,
    command=view_student_profile
).pack(
    side="left",
    padx=5
)

tk.Button(
    table_button_frame,
    text="✏ Edit Selected",
    font=get_font(10, True),
    bg="#f59e0b",
    fg="white",
    padx=15,
    pady=7,
    command=edit_selected_student
).pack(
    side="left",
    padx=5
)

tk.Button(
    table_button_frame,
    text="🗑 Delete Selected",
    font=get_font(10, True),
    bg="#dc2626",
    fg="white",
    padx=15,
    pady=7,
    command=delete_selected_student
).pack(
    side="left",
    padx=5
)

tk.Button(
    table_button_frame,
    text="🎫 Generate ID Card",
    font=get_font(10, True),
    bg="#0891b2",
    fg="white",
    padx=15,
    pady=7,
    command=generate_id_card
).pack(
    side="left",
    padx=5
)

tk.Button(
    table_button_frame,
    text="📊 Attendance",
    font=get_font(10, True),
    bg="#4f46e5",
    fg="white",
    padx=15,
    pady=7,
    command=open_attendance_window
).pack(
    side="left",
    padx=5
)

tk.Button(
    table_button_frame,
    text="📅 Monthly Report",
    font=get_font(10, True),
    bg="#9333ea",
    fg="white",
    padx=15,
    pady=7,
    command=open_monthly_report
).pack(
    side="left",
    padx=5
)

tk.Button(
    table_button_frame,
    text="📥 Export Students",
    font=get_font(10, True),
    bg="#0284c7",
    fg="white",
    padx=15,
    pady=7,
    command=export_to_excel
).pack(
    side="left",
    padx=5
)


# =========================================================
# LOGOUT BUTTON
# =========================================================

tk.Button(
    main_frame,
    text="🚪 Logout",
    font=("Arial", 11, "bold"),
    bg="#374151",
    fg="white",
    padx=25,
    pady=9,
    command=logout
).pack(
    pady=25
)


# =========================================================
# LOGIN WINDOW
# =========================================================

login_window = tk.Toplevel(window)

login_window.title(
    "Login - Student Management System"
)

login_window.geometry(
    "450x400"
)

login_window.resizable(
    False,
    False
)

login_window.configure(
    bg="#1f3b56"
)


tk.Label(
    login_window,
    text="Student Management System",
    font=("Arial", 22, "bold"),
    bg="#1f3b56",
    fg="white"
).pack(
    pady=(45, 10)
)

tk.Label(
    login_window,
    text="Admin Login",
    font=("Arial", 14),
    bg="#1f3b56",
    fg="#dbeafe"
).pack(
    pady=(0, 25)
)


login_form = tk.Frame(
    login_window,
    bg="#1f3b56"
)

login_form.pack()


tk.Label(
    login_form,
    text="Username",
    font=get_font(11, True),
    bg="#1f3b56",
    fg="white"
).grid(
    row=0,
    column=0,
    sticky="w",
    pady=8
)

username_login_entry = tk.Entry(
    login_form,
    font=get_font(12),
    width=25
)

username_login_entry.grid(
    row=1,
    column=0,
    pady=(0, 15)
)


tk.Label(
    login_form,
    text="Password",
    font=get_font(11, True),
    bg="#1f3b56",
    fg="white"
).grid(
    row=2,
    column=0,
    sticky="w",
    pady=8
)

password_login_entry = tk.Entry(
    login_form,
    font=get_font(12),
    width=25,
    show="*"
)

password_login_entry.grid(
    row=3,
    column=0,
    pady=(0, 5)
)


show_password_var = tk.BooleanVar(
    value=False
)

tk.Checkbutton(
    login_form,
    text="Show Password",
    variable=show_password_var,
    command=toggle_password,
    bg="#1f3b56",
    fg="white",
    selectcolor="#1f3b56",
    activebackground="#1f3b56",
    activeforeground="white"
).grid(
    row=4,
    column=0,
    pady=5
)


tk.Button(
    login_window,
    text="Login",
    font=("Arial", 12, "bold"),
    bg="#2563eb",
    fg="white",
    padx=40,
    pady=10,
    command=login
).pack(
    pady=20
)


tk.Label(
    login_window,
    text="Username: admin    Password: 1234",
    font=("Arial", 9),
    bg="#1f3b56",
    fg="#bfdbfe"
).pack(
    pady=5
)


# Enter key login
password_login_entry.bind(
    "<Return>",
    lambda event: login()
)

username_login_entry.bind(
    "<Return>",
    lambda event: login()
)


# =========================================================
# INITIAL LOAD
# =========================================================

load_students()
load_course_filter()
update_dashboard()


# Hide main window until login
window.withdraw()

username_login_entry.focus()


# =========================================================
# MOUSE WHEEL SCROLL
# =========================================================

def mouse_wheel(event):

    main_canvas.yview_scroll(
        int(-1 * (event.delta / 120)),
        "units"
    )


main_canvas.bind_all(
    "<MouseWheel>",
    mouse_wheel
)


# =========================================================
# START APPLICATION
# =========================================================

window.mainloop()